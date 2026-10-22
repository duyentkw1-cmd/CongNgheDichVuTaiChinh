import csv
import io
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

from .db import now
from .domain import CATEGORIES, REVENUE_CATEGORIES, summarize

HEADERS = ['Ngày', 'Nội dung', 'Nhóm', 'Thu (VND)', 'Chi (VND)', 'Thanh toán', 'Phạm vi']

def safe_text(v):
    # Spreadsheet injection protection; data from users must remain literal strings.
    return "'" + v if isinstance(v, str) and v.lstrip().startswith(('=', '+', '-', '@')) else v

def export_file(rows, start, end, household, kind, fmt, scope='business'):
    if kind == 'revenue':
        rows = [t for t in rows if t.kind == 'income' and t.category in REVENUE_CATEGORIES and t.scope == 'business']
    summary = summarize(rows, start, end)
    title = 'SỔ DOANH THU BÁN HÀNG VÀ DỊCH VỤ' if kind == 'revenue' else 'SỔ THEO DÕI THU – CHI'
    scope_label = {'all': 'Tất cả', 'business': 'Kinh doanh', 'personal': 'Cá nhân'}[scope]
    metadata = [household, f'Kỳ: {start:%d/%m/%Y} – {end:%d/%m/%Y} | Phạm vi: {scope_label}',
                f'Xuất lúc: {now():%d/%m/%Y %H:%M} (Asia/Ho_Chi_Minh)',
                'Báo cáo quản lý nội bộ từ giao dịch đã xác nhận; chênh lệch thu – chi không phải lợi nhuận.']
    records = [[t.date.strftime('%d/%m/%Y'), safe_text(t.description), CATEGORIES[t.category][0],
                int(t.amount) if t.kind == 'income' else 0, int(t.amount) if t.kind == 'expense' else 0,
                'Tiền mặt' if t.payment == 'cash' else 'Chuyển khoản',
                'Kinh doanh' if t.scope == 'business' else 'Cá nhân'] for t in reversed(rows)]
    totals = ['', 'TỔNG CỘNG', '', summary['income'], summary['expense'], '', '']
    if fmt == 'csv':
        text = io.StringIO(newline='')
        writer = csv.writer(text)
        writer.writerow([title])
        for m in metadata:
            writer.writerow([safe_text(m)])
        writer.writerow(HEADERS)
        writer.writerows(records)
        writer.writerow(totals)
        return ('\ufeff' + text.getvalue()).encode(), 'text/csv; charset=utf-8'
    if fmt == 'xlsx':
        book = Workbook()
        sheet = book.active
        sheet.title = 'Doanh thu' if kind == 'revenue' else 'Thu chi'
        sheet.append([title])
        for m in metadata:
            sheet.append([safe_text(m)])
        sheet.append(HEADERS)
        for row in records + [totals]:
            sheet.append(row)
        for row in (sheet[1], sheet[6], sheet[sheet.max_row]):
            for cell in row:
                cell.fill = PatternFill('solid', fgColor='155D4C')
                cell.font = Font(color='FFFFFF', bold=True)
        for col, width in zip('ABCDEFG', [15, 42, 30, 19, 19, 21, 18]):
            sheet.column_dimensions[col].width = width
        for row in sheet.iter_rows(min_row=7, min_col=4, max_col=5):
            for cell in row:
                cell.number_format = '#,##0'
        sheet.freeze_panes = 'A7'
        sheet.auto_filter.ref = f'A6:G{max(6, sheet.max_row - 1)}'
        output = io.BytesIO()
        book.save(output)
        return output.getvalue(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    font_dir = Path(__file__).parent / 'assets'
    for name, file in [('VN', 'DejaVuSans.ttf'), ('VN-Bold', 'DejaVuSans-Bold.ttf')]:
        if name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(name, str(font_dir / file)))
    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=landscape(A4), rightMargin=28, leftMargin=28, topMargin=28, bottomMargin=28)
    body = ParagraphStyle('body', fontName='VN', fontSize=8, leading=12)
    heading = ParagraphStyle('heading', fontName='VN-Bold', fontSize=15, leading=22, textColor=colors.HexColor('#155D4C'))
    elements = [Paragraph(title, heading)] + [Paragraph(escape(m), body) for m in metadata] + [Spacer(1, 16)]
    table_rows = [HEADERS]
    for row in records + [totals]:
        table_rows.append([f'{v:,}'.replace(',', '.') if isinstance(v, int) else Paragraph(escape(str(v)), body) for v in row])
    table = Table(table_rows, colWidths=[66, 200, 150, 88, 88, 95, 98], repeatRows=1)
    table.setStyle(TableStyle([('FONTNAME', (0, 0), (-1, -1), 'VN'), ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#DDEDE4')),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#DDEDE4')),
        ('GRID', (0, 0), (-1, -1), .3, colors.HexColor('#D7DFDB')),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('BOTTOMPADDING', (0, 0), (-1, -1), 7),
        ('TOPPADDING', (0, 0), (-1, -1), 7)]))
    elements.append(table)
    def footer(canvas, document):
        canvas.setFont('VN', 8)
        canvas.drawRightString(810, 14, f'Trang {document.page}')
    doc.build(elements, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue(), 'application/pdf'
