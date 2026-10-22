"""Focused integration tests for money, boundaries, authentication and export consistency."""
import csv
import io
import os
import tempfile
import unittest
import uuid
from datetime import timedelta
from pathlib import Path
from unittest.mock import patch

from _isolation import TEMP_DIR

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from backend.main import app
from backend.db import SessionLocal, Household, User, hash_password, today
from backend.local_parse import parse_fields

HEADERS = {'X-Requested-With': 'so-thu-chi'}

class IntegrationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.client = TestClient(app, headers=HEADERS)
        cls.client.__enter__()
        r = cls.client.post('/api/auth/login', json={'email': 'demo@sothuchi.vn', 'password': 'Demo123!'})
        assert r.status_code == 200, r.text

    @classmethod
    def tearDownClass(cls):
        cls.client.__exit__(None, None, None)

    def payload(self, **values):
        return {'kind':'expense','amount':123456,'description':'Kiểm thử '+str(uuid.uuid4()),
                'category':'shipping','date':today().isoformat(),'payment':'cash','scope':'business',
                'source':'manual','request_id':str(uuid.uuid4()),**values}

    def test_01_auth_and_csrf(self):
        with TestClient(app) as anonymous:
            self.assertEqual(anonymous.get('/api/report').status_code,401)
            self.assertEqual(anonymous.get('/api/exports?start=2026-01-01&end=2026-01-31').status_code,401)
            self.assertEqual(anonymous.post('/api/transactions',json=self.payload()).status_code,403)
            r=anonymous.post('/api/auth/login',headers=HEADERS,json={'email':'demo@sothuchi.vn','password':'wrong'})
            self.assertEqual(r.status_code,401)

    def test_02_idempotency_and_duplicate_warning(self):
        p=self.payload()
        a=self.client.post('/api/transactions',json=p)
        self.assertEqual(a.status_code,200,a.text)
        b=self.client.post('/api/transactions',json=p)
        self.assertEqual(a.json()['id'],b.json()['id'])
        c=self.client.post('/api/transactions',json={**p,'request_id':str(uuid.uuid4())})
        self.assertEqual(c.status_code,409)
        self.assertEqual(self.client.post('/api/transactions',json={**p,'request_id':str(uuid.uuid4()),'allow_duplicate':True}).status_code,200)

    def test_03_money_validation(self):
        for val in [-1,0,1.5,1000000000001]:
            self.assertEqual(self.client.post('/api/transactions',json=self.payload(amount=val)).status_code,422)
        self.assertEqual(self.client.post('/api/transactions',json=self.payload(category='sales')).status_code,422)
        self.assertEqual(self.client.post('/api/transactions',json=self.payload(date=(today()+timedelta(days=1)).isoformat())).status_code,422)

    def test_04_export_matches_report(self):
        report=self.client.get('/api/report?period=month').json()
        query=f"start={report['start']}&end={report['end']}&scope=business&kind=all"
        r=self.client.get('/api/exports?'+query+'&format=xlsx')
        self.assertEqual(r.status_code,200,r.text[:100])
        sheet=load_workbook(io.BytesIO(r.content),data_only=True).active
        self.assertEqual(sheet.cell(sheet.max_row,4).value,report['income'])
        self.assertEqual(sheet.cell(sheet.max_row,5).value,report['expense'])
        self.assertEqual(sum(sheet.cell(i,4).value for i in range(7,sheet.max_row)),report['income'])
        rows=list(csv.reader(io.StringIO(self.client.get('/api/exports?'+query+'&format=csv').content.decode('utf-8-sig'))))
        self.assertEqual(int(rows[-1][3]),report['income'])
        self.assertEqual(int(rows[-1][4]),report['expense'])
        r=self.client.get('/api/exports?'+query+'&format=pdf')
        self.assertEqual(r.status_code,200)
        self.assertTrue(r.content.startswith(b'%PDF'))
        Path('qa').mkdir(exist_ok=True)
        Path('qa/report.pdf').write_bytes(r.content)

    def test_05_loan_not_revenue_and_personal_scope(self):
        initial=self.client.get('/api/report?period=day').json()
        self.client.post('/api/transactions',json=self.payload(kind='income',category='capital',amount=1000000))
        after=self.client.get('/api/report?period=day').json()
        self.assertEqual(after['revenue'],initial['revenue'])
        self.assertEqual(after['income']-initial['income'],1000000)
        res=self.client.post('/api/transactions',json=self.payload(scope='business',category='personal',amount=10000))
        self.assertEqual(res.status_code,422)

    def test_06_budget_exact_boundary(self):
        report=self.client.get('/api/report?period=day').json()
        actual=sum(t['amount'] for t in report['transactions'] if t['category']=='shipping' and t['kind']=='expense')
        for limit,expected in [(actual+1,False),(actual,False),(actual-1,True)]:
            r=self.client.post('/api/budgets',json={'category':'shipping','period':'day','amount':limit})
            row=next(x for x in r.json() if x['category']=='shipping' and x['period']=='day')
            self.assertEqual(row['exceeded'],expected)

    def test_07_update_delete_and_isolation(self):
        p=self.payload();r=self.client.post('/api/transactions',json=p);tid=r.json()['id']
        r=self.client.put('/api/transactions/'+tid,json={**p,'amount':765432})
        self.assertEqual(r.json()['amount'],765432)
        self.assertEqual(len(self.client.get('/api/transactions/'+tid+'/history').json()),2)
        with SessionLocal() as db:
            db.add(Household(id='other',name='Other'));db.flush()
            db.add(User(id='other-user',household_id='other',name='Other',email='other@test.local',password_hash=hash_password('test-pass')));db.commit()
        with TestClient(app,headers=HEADERS) as other:
            other.post('/api/auth/login',json={'email':'other@test.local','password':'test-pass'})
            self.assertEqual(other.get('/api/report').json()['count'],0)
            self.assertEqual(other.put('/api/transactions/'+tid,json=p).status_code,404)
            self.assertEqual(other.delete('/api/transactions/'+tid).status_code,404)
        self.assertEqual(self.client.delete('/api/transactions/'+tid).status_code,200)
        self.assertNotIn(tid,[t['id'] for t in self.client.get('/api/report?period=day').json()['transactions']])

    def test_08_dates_and_parser(self):
        for text,amount in [('mua nguyên liệu 800 nghìn',800000),('thu tiền bán hàng 2 triệu',2000000),
                            ('Mua nguyên liệu hết sáu trăm nghìn',600000),('trả tiền điện một triệu rưỡi',1500000),
                            ('Tổng cộng: 650.000 VND',650000)]:
            self.assertEqual(parse_fields(text)['amount'],amount,text)
        self.assertIsNone(parse_fields('mua rau')['amount'])
        self.assertEqual(self.client.get('/api/report?start=2026-12-01&end=2026-01-01').status_code,422)
        r=self.client.get('/api/report?period=week').json()
        from datetime import date
        self.assertEqual(date.fromisoformat(r['start']).weekday(),0)
        self.assertEqual((date.fromisoformat(r['end'])-date.fromisoformat(r['start'])).days,6)

    def test_09_upload_rejection_and_no_fake_ocr(self):
        r=self.client.post('/api/attachments',files={'file':('bad.png',b'not an image','image/png')})
        self.assertEqual(r.status_code,400)
        r=self.client.post('/api/ai/image',files={'file':('sample.png',b'not an image','image/png')})
        self.assertIn(r.status_code,(400,503))
        r=self.client.post('/api/ai/text',json={'text':'mua nguyên liệu 300 nghìn'})
        self.assertEqual(r.json()['engine'],'local')
        r=self.client.post('/api/transactions',json=self.payload(description='=HYPERLINK("x")'))
        report=self.client.get('/api/report?period=day').json()
        result=self.client.get('/api/exports',params={'start':report['start'],'end':report['end'],'format':'csv'})
        self.assertIn("'=HYPERLINK",result.content.decode('utf-8-sig'))

    def test_10_local_model_and_dataset(self):
        r=self.client.post('/api/ai/text',json={'text':'Thu tiền bán hàng 2 triệu, chuyển khoản'})
        self.assertEqual(r.status_code,200,r.text)
        self.assertEqual(r.json()['engine'],'local')
        self.assertEqual(r.json()['amount'],2000000)
        self.assertEqual(r.json()['category'],'sales')
        self.assertIn('classification',r.json())
        r=self.client.get('/api/ai/dataset/train')
        self.assertEqual(r.status_code,200)
        rows=list(csv.DictReader(io.StringIO(r.content.decode('utf-8-sig'))))
        self.assertGreaterEqual(len(rows),1056)
        self.assertEqual(self.client.get('/api/ai/dataset/other').status_code,422)

    def test_11_ai_status_requires_login_and_ignores_external_keys(self):
        with TestClient(app) as anonymous:
            self.assertEqual(anonymous.get('/api/ai/status').status_code,401)
            self.assertEqual(anonymous.get('/api/ai/dataset/train').status_code,401)
        with patch.dict(os.environ,{'AI_MODE':'live','GEMINI_API_KEY':'test-only-secret-key'}):
            r=self.client.get('/api/ai/status')
            self.assertEqual(r.status_code,200)
            self.assertTrue(r.json()['text'])
            self.assertEqual(r.json()['mode'],'local')
            self.assertFalse(r.json()['network_required'])
            self.assertNotIn('test-only-secret-key',r.text)

if __name__=='__main__':
    unittest.main(verbosity=2)
